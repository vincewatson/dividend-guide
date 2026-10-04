#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""博客（子弹列车公众号文章目录）数据生成 —— 手动维护，不参与自动流水线。
============================================================
输入：user_upload/公众号历史文章(20240123-20260919).xlsx
      （列：发表日期 / 标题 / 文章链接 / 所属栏目）
输出：data/blogData.json
      字段：date / title / url / column / direction / indexes[] / paid

说明（2026-10-04 用户上传「全部历史文章」后重写）：
  - 该表为公众号全部历史文章的权威清单（289 篇，2024-01-23 ~ 2026-09-19），
    含「所属栏目」字段；
  - 「内容标签」（direction）与「相关指数」（indexes）本表暂缺 → 由
    data/blogAnnotations.json 逐条补录（保留用户此前给出的标注），未标注的留空，
    待日后接入后台再补；
  - 「付费文章」为作者标注的所属栏目之一（2 篇）→ paid=true。
  - 按链接去重、按发表日期倒序；应用全站空格规范。

用法：python3 sync_blog.py   （生成后需 python3 embed_data.py 刷新内嵌兜底）
"""
import glob, io, json, os, re
import openpyxl

BASE = os.path.dirname(os.path.abspath(__file__))
SRC = sorted(glob.glob(os.path.join(BASE, "user_upload/公众号历史文章*.xlsx")))[-1]
OUT = os.path.join(BASE, "data/blogData.json")
ANNOT_PATH = os.path.join(BASE, "data/blogAnnotations.json")   # url -> {direction, indexes}（内容标签/相关指数，逐条补录）

# 「付费文章」作为所属栏目值时，同时标记 paid（前端「文章属性」筛选承载）
PAID_COL = "付费文章"

CJK = '\u3400-\u4dbf\u4e00-\u9fff\uf900-\ufaff'
R1 = re.compile(r'(?<=[' + CJK + r']) +(?=[0-9A-Za-z])')
R2 = re.compile(r'(?<=[0-9A-Za-z]) +(?=[' + CJK + r'])')


def norm(s):
    return R2.sub('', R1.sub('', s)) if isinstance(s, str) else s


def main():
    wb = openpyxl.load_workbook(SRC, data_only=True)
    ws = wb.active
    recs = []
    for r in range(2, ws.max_row + 1):
        d = ws.cell(r, 1).value
        title = (ws.cell(r, 2).value or '')
        url = (ws.cell(r, 3).value or '')
        col = ws.cell(r, 4).value
        title = title.strip() if isinstance(title, str) else str(title)
        url = url.strip() if isinstance(url, str) else str(url)
        col = (col.strip() if isinstance(col, str) else (str(col) if col else ''))
        if not title and not url and not d:
            continue
        recs.append({
            "date": d.strftime("%Y-%m-%d") if hasattr(d, "strftime") else (str(d).strip() if d else ""),
            "title": norm(title),
            "url": url,
            "column": norm(col),
            "direction": "",      # 内容标签：本表缺 → 留空
            "indexes": [],        # 相关指数：本表缺 → 留空
        })

    # 按链接去重（同一篇只保留一条；合并日期/栏目）
    by_url = {}
    for r in recs:
        k = r["url"] or (r["date"] + "|" + r["title"])
        if k not in by_url:
            by_url[k] = r
        else:
            a = by_url[k]
            if not a["column"] and r["column"]:
                a["column"] = r["column"]
            if not a["title"] and r["title"]:
                a["title"] = r["title"]
            if r["date"] and (not a["date"] or r["date"] > a["date"]):
                a["date"] = r["date"]

    # 补录：内容标签（direction）/ 相关指数（indexes）——由 data/blogAnnotations.json 按 url 合并
    annot = {}
    if os.path.exists(ANNOT_PATH):
        with io.open(ANNOT_PATH, encoding="utf-8") as f:
            annot = json.load(f)
    for r in by_url.values():
        a = annot.get(r["url"])
        if a:
            if a.get("direction"):
                r["direction"] = norm(a["direction"])
            if a.get("indexes"):
                r["indexes"] = [norm(x) for x in a["indexes"]]

    out = sorted(by_url.values(), key=lambda x: x["date"], reverse=True)
    for r in out:
        r["paid"] = (r["column"] == PAID_COL)

    with io.open(OUT, "w", encoding="utf-8") as f:
        json.dump(out, f, ensure_ascii=False)

    from collections import Counter
    print("来源:", os.path.basename(SRC))
    print("文章:", len(out), "| 付费:", sum(1 for x in out if x["paid"]),
          "| 有内容标签:", sum(1 for x in out if x["direction"]),
          "| 有相关指数:", sum(1 for x in out if x["indexes"]))
    print("年份:", dict(sorted(Counter(x["date"][:4] for x in out).items())))
    print("所属栏目:", dict(Counter(x["column"] for x in out).most_common()))
    print("写入", OUT, os.path.getsize(OUT), "bytes")


if __name__ == "__main__":
    main()
